"""
engine.py
---------
All Excel processing logic lives here.

Responsibilities:
    - Read Source workbook
    - Select sheets according to workspace
    - Exclude Instruction, Instructions and Data Validation
    - Copy headers and formatting
    - Create missing sheets
    - Clear old data
    - Write Source data into Template
    - Handle special L4 Task Attributes sheet
    - Return updated Template workbook

No Streamlit dependency.
"""

from copy import copy
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell


# ============================================================
# CUSTOM ERROR
# ============================================================

class ProcessingError(Exception):
    """Raised for expected, user-facing processing failures."""

    pass


# ============================================================
# CONSTANTS
# ============================================================

EXCLUDED_SHEETS = {
    "instruction",
    "instructions",
    "data validation",
}

# Special L4 Task Attributes sheet
SPECIAL_L4_TASK_ATTRIBUTES_SHEET = "l4 task attributes"


# ============================================================
# SHEET HELPERS
# ============================================================

def normalize_sheet_name(sheet_name):
    """
    Normalize a sheet name for comparison.

    Example:
        ' L4 Task Attributes ' -> 'l4 task attributes'
        'INSTRUCTIONS' -> 'instructions'
    """

    return str(sheet_name).strip().lower()


def is_excluded_sheet(sheet_name):
    """
    Returns True if the sheet should never be processed.

    Excluded sheets:
        - Instruction
        - Instructions
        - Data Validation

    Comparison is case-insensitive and ignores
    leading/trailing spaces.
    """

    return (
        normalize_sheet_name(sheet_name)
        in EXCLUDED_SHEETS
    )


def is_l4_task_attributes_sheet(sheet_name):
    """
    Returns True when the sheet name is exactly:

        L4 Task Attributes

    Comparison is case-insensitive and ignores
    leading/trailing spaces.
    """

    return (
        normalize_sheet_name(sheet_name)
        == SPECIAL_L4_TASK_ATTRIBUTES_SHEET
    )


# ============================================================
# L4 FLOW DIAGRAM SHEET HELPER
# ============================================================

def starts_with_alpha_and_has_digit(sheet_name):
    """
    Used by L4 Flow Diagrams instead of a fixed 'PP-' prefix check.

    Returns True when the sheet name:
        - Starts with an alphabet letter (A-Z / a-z)
        - Contains at least one digit somewhere AFTER that first
          letter

    Examples:
        'PP-100'      -> True
        'Flow2'       -> True
        'Diagram A1'  -> True
        'Summary'     -> False
        '100-Flow'    -> False
        ''            -> False
    """

    name = str(sheet_name).strip()

    if not name:
        return False

    if not name[0].isalpha():
        return False

    return any(
        character.isdigit()
        for character in name[1:]
    )


# ============================================================
# DETERMINE DATA START ROW
# ============================================================

def get_data_start_row(
    sheet_name,
    mode
):
    """
    Determine where data should start in the Source sheet.

    Normal behavior:
        Data starts from Row 3.

    Special behavior:
        For L4 Task Attributes sheet named
        'L4 Task Attributes', data starts from Row 2.
    """

    if (
        mode == "task_attributes"
        and is_l4_task_attributes_sheet(sheet_name)
    ):
        return 2

    return 3


# ============================================================
# CELL COPYING
# ============================================================

def copy_cell(
    source_cell,
    target_cell
):
    """
    Copy value and complete formatting safely.

    Handles MergedCell safely.
    """

    if isinstance(
        target_cell,
        MergedCell
    ):
        return

    # --------------------------------------------------------
    # VALUE
    # --------------------------------------------------------

    target_cell.value = source_cell.value

    # --------------------------------------------------------
    # STYLE
    # --------------------------------------------------------

    if source_cell.has_style:

        target_cell._style = copy(
            source_cell._style
        )

    target_cell.font = copy(
        source_cell.font
    )

    target_cell.fill = copy(
        source_cell.fill
    )

    target_cell.border = copy(
        source_cell.border
    )

    target_cell.alignment = copy(
        source_cell.alignment
    )

    target_cell.number_format = (
        source_cell.number_format
    )

    target_cell.protection = copy(
        source_cell.protection
    )

    # --------------------------------------------------------
    # HYPERLINK
    # --------------------------------------------------------

    if source_cell.hyperlink:

        target_cell._hyperlink = copy(
            source_cell.hyperlink
        )

    # --------------------------------------------------------
    # COMMENT
    # --------------------------------------------------------

    if source_cell.comment:

        target_cell.comment = copy(
            source_cell.comment
        )


# ============================================================
# HEADER COPYING
# ============================================================

def copy_header(
    source_ws,
    target_ws
):
    """
    Copy the first two rows from Source to Target.

    Also copies:
        - Formatting
        - Row heights
        - Column widths
        - Hidden rows/columns
        - Merged cells
    """

    # --------------------------------------------------------
    # COPY FIRST TWO ROWS
    # --------------------------------------------------------

    for row in range(1, 3):

        target_ws.row_dimensions[
            row
        ].height = (
            source_ws.row_dimensions[
                row
            ].height
        )

        target_ws.row_dimensions[
            row
        ].hidden = (
            source_ws.row_dimensions[
                row
            ].hidden
        )

        for col in range(
            1,
            source_ws.max_column + 1
        ):

            source_cell = (
                source_ws.cell(
                    row=row,
                    column=col
                )
            )

            target_cell = (
                target_ws.cell(
                    row=row,
                    column=col
                )
            )

            copy_cell(
                source_cell,
                target_cell
            )

    # --------------------------------------------------------
    # COPY COLUMN WIDTHS
    # --------------------------------------------------------

    for (
        col_letter,
        dimension
    ) in source_ws.column_dimensions.items():

        target_ws.column_dimensions[
            col_letter
        ].width = dimension.width

        target_ws.column_dimensions[
            col_letter
        ].hidden = dimension.hidden

    # --------------------------------------------------------
    # COPY MERGED CELLS
    # --------------------------------------------------------

    existing_merges = {
        str(x)
        for x in target_ws.merged_cells.ranges
    }

    for merged_range in (
        source_ws.merged_cells.ranges
    ):

        merge_string = str(
            merged_range
        )

        if merge_string in existing_merges:
            continue

        try:

            target_ws.merge_cells(
                merge_string
            )

        except Exception:
            pass


# ============================================================
# READ SOURCE DATA
# ============================================================

def read_sheet_data(
    source_excel_file,
    sheet,
    start_row=3
):
    """
    Read data from the specified starting row.

    Examples:

        start_row=3
            Reads Row 3 onwards.

        start_row=2
            Reads Row 2 onwards.

    pandas skiprows is zero-based, therefore:

        Row 3 -> skiprows=2
        Row 2 -> skiprows=1
    """

    skip_rows = start_row - 1

    try:

        df = pd.read_excel(
            source_excel_file,
            sheet_name=sheet,
            skiprows=skip_rows,
            header=None,
            dtype=str,
        )

    except Exception as e:

        raise ProcessingError(
            f"Could not read sheet '{sheet}': {e}"
        )

    # --------------------------------------------------------
    # REMOVE COMPLETELY EMPTY ROWS
    # --------------------------------------------------------

    df = (
        df
        .dropna(how="all")
        .reset_index(drop=True)
    )

    # --------------------------------------------------------
    # REPLACE NaN
    # --------------------------------------------------------

    df = df.fillna("")

    # --------------------------------------------------------
    # CONVERT TO STRING
    # --------------------------------------------------------

    df = df.astype(str)

    return df


# ============================================================
# CLEAR OLD DATA
# ============================================================

def clear_old_data(
    ws,
    start_row=3
):
    """
    Delete existing data from start_row onwards.

    Examples:

        start_row=3
            Keeps Row 1 and Row 2.

        start_row=2
            Keeps only Row 1.
    """

    if ws.max_row >= start_row:

        ws.delete_rows(
            start_row,
            ws.max_row - start_row + 1
        )


# ============================================================
# PREPARE TARGET SHEET
# ============================================================

def prepare_target_sheet(
    source_ws,
    target_wb,
    sheet_name,
    start_row=3
):
    """
    Prepare destination sheet.

    Existing sheet:
        - Clear old data from start_row
        - Keep rows above start_row

    Missing sheet:
        - Create sheet
        - Copy first two header rows
        - Copy formatting
    """

    # --------------------------------------------------------
    # EXISTING SHEET
    # --------------------------------------------------------

    if sheet_name in target_wb.sheetnames:

        target_ws = target_wb[
            sheet_name
        ]

        clear_old_data(
            target_ws,
            start_row
        )

        return target_ws

    # --------------------------------------------------------
    # CREATE NEW SHEET
    # --------------------------------------------------------

    target_ws = target_wb.create_sheet(
        title=sheet_name
    )

    copy_header(
        source_ws,
        target_ws
    )

    return target_ws


# ============================================================
# WRITE DATA
# ============================================================

def write_dataframe(
    target_ws,
    df,
    start_row=3
):
    """
    Write DataFrame starting from start_row.
    """

    for row_index, row_data in enumerate(
        df.itertuples(
            index=False,
            name=None
        ),
        start=start_row
    ):

        for column_index, value in enumerate(
            row_data,
            start=1
        ):

            target_ws.cell(
                row=row_index,
                column=column_index
            ).value = value


# ============================================================
# SHEET SELECTION
# ============================================================

def select_sheets(
    mode,
    sheet_names
):
    """
    Select sheets according to workspace.

    Modes:

        all
            All sheets except:
                Instruction
                Instructions
                Data Validation

        pp
            Sheets whose name:
                - Starts with a letter
                - Contains a number somewhere after
                  the first character

            Examples:
                PP-100
                Flow2
                Diagram A1

        task_attributes
            All sheets except:
                Instruction
                Instructions
                Data Validation

        user_stories
            Only User Stories.

    IMPORTANT:
        Instruction, Instructions and Data Validation
        are ALWAYS excluded.
    """

    # --------------------------------------------------------
    # REMOVE EXCLUDED SHEETS
    # --------------------------------------------------------

    available_sheets = [
        sheet
        for sheet in sheet_names
        if not is_excluded_sheet(sheet)
    ]

    # --------------------------------------------------------
    # L0-L3
    # --------------------------------------------------------

    if mode == "all":

        return available_sheets

    # --------------------------------------------------------
    # L4 FLOW DIAGRAMS
    # --------------------------------------------------------

    if mode == "pp":

        return [
            sheet
            for sheet in available_sheets
            if (
                not is_excluded_sheet(sheet)
                and starts_with_alpha_and_has_digit(sheet)
            )
        ]

    # --------------------------------------------------------
    # L4 TASK ATTRIBUTES
    # --------------------------------------------------------

    if mode == "task_attributes":

        # Instructions / Instruction and Data Validation
        # have already been removed from available_sheets.
        #
        # The ONLY special behavior is that if the
        # actual sheet name is "L4 Task Attributes",
        # its data starts from Row 2.

        return available_sheets

    # --------------------------------------------------------
    # USER STORIES
    # --------------------------------------------------------

    if mode == "user_stories":

        for sheet in available_sheets:

            if normalize_sheet_name(
                sheet
            ) == "user stories":

                return [sheet]

        return []

    # --------------------------------------------------------
    # INVALID MODE
    # --------------------------------------------------------

    raise ValueError(
        f"Unknown mode: {mode}"
    )


# ============================================================
# SOURCE SHEET PREVIEW
# ============================================================

def list_source_sheets(
    source_bytes
):
    """
    Return source workbook sheet names.
    """

    try:

        return pd.ExcelFile(
            BytesIO(source_bytes)
        ).sheet_names

    except Exception:

        return None


# ============================================================
# MAIN PROCESSING FUNCTION
# ============================================================

def process_workbook(
    source_bytes: bytes,
    target_bytes: bytes,
    mode: str,
    progress_callback=None
):
    """
    Process Source workbook and update Template workbook.

    Normal sheets:
        Read from Row 3.

    Special L4 Task Attributes sheet:
        If sheet name is 'L4 Task Attributes',
        read from Row 2.

    Excluded:
        Instruction
        Instructions
        Data Validation
    """

    log_lines = []

    # ========================================================
    # OPEN SOURCE WITH PANDAS
    # ========================================================

    try:

        source_excel = pd.ExcelFile(
            BytesIO(source_bytes)
        )

    except Exception as e:

        raise ProcessingError(
            "Could not open the Source Excel file: "
            f"{e}"
        )

    # ========================================================
    # OPEN SOURCE WITH OPENPYXL
    # ========================================================

    try:

        source_wb = load_workbook(
            BytesIO(source_bytes)
        )

    except Exception as e:

        raise ProcessingError(
            "Could not open the Source Excel file "
            f"for formatting: {e}"
        )

    # ========================================================
    # OPEN TEMPLATE
    # ========================================================

    try:

        target_wb = load_workbook(
            BytesIO(target_bytes)
        )

    except Exception as e:

        raise ProcessingError(
            "Could not open the Template Excel file: "
            f"{e}"
        )

    # ========================================================
    # SOURCE SHEETS
    # ========================================================

    source_sheet_names = (
        source_excel.sheet_names
    )

    # ========================================================
    # SELECT SHEETS
    # ========================================================

    sheets = select_sheets(
        mode,
        source_sheet_names
    )

    # ========================================================
    # NO SHEETS
    # ========================================================

    if not sheets:

        if mode == "pp":

            raise ProcessingError(
                "No matching sheets were found. A sheet must "
                "start with a letter and contain a number "
                "elsewhere in its name (for example 'PP-100' "
                "or 'Flow2')."
            )

        elif mode == "user_stories":

            raise ProcessingError(
                "The 'User Stories' sheet was not "
                "found in the Source workbook."
            )

        else:

            raise ProcessingError(
                "No sheets are available for processing "
                "after excluding 'Instruction', "
                "'Instructions' and 'Data Validation'."
            )

    # ========================================================
    # LOG
    # ========================================================

    log_lines.append(
        f"{len(sheets)} sheet(s) detected for processing."
    )

    log_lines.append(
        "Skipped sheets: "
        "'Instruction', 'Instructions' and "
        "'Data Validation'."
    )

    # ========================================================
    # PROCESS SHEETS
    # ========================================================

    for index, sheet_name in enumerate(
        sheets
    ):

        # ----------------------------------------------------
        # SAFETY CHECK
        # ----------------------------------------------------

        if is_excluded_sheet(
            sheet_name
        ):
            continue

        # ----------------------------------------------------
        # PROGRESS
        # ----------------------------------------------------

        if progress_callback:

            progress_callback(
                index,
                len(sheets),
                sheet_name
            )

        # ----------------------------------------------------
        # SOURCE WORKSHEET
        # ----------------------------------------------------

        try:

            source_ws = source_wb[
                sheet_name
            ]

        except KeyError:

            raise ProcessingError(
                f"Could not find source sheet "
                f"'{sheet_name}'."
            )

        # ----------------------------------------------------
        # DETERMINE START ROW
        # ----------------------------------------------------

        data_start_row = get_data_start_row(
            sheet_name,
            mode
        )

        # ----------------------------------------------------
        # READ DATA
        # ----------------------------------------------------

        df = read_sheet_data(
            source_excel,
            sheet_name,
            start_row=data_start_row
        )

        # ----------------------------------------------------
        # PREPARE TARGET
        # ----------------------------------------------------

        target_ws = prepare_target_sheet(
            source_ws,
            target_wb,
            sheet_name,
            start_row=data_start_row
        )

        # ----------------------------------------------------
        # WRITE DATA
        # ----------------------------------------------------

        write_dataframe(
            target_ws,
            df,
            start_row=data_start_row
        )

        # ----------------------------------------------------
        # LOG SPECIAL CASE
        # ----------------------------------------------------

        if (
            mode == "task_attributes"
            and is_l4_task_attributes_sheet(
                sheet_name
            )
        ):

            log_lines.append(
                f"Processed '{sheet_name}' "
                f"from Row 2 "
                f"({len(df)} data row(s))."
            )

        else:

            log_lines.append(
                f"Processed '{sheet_name}' "
                f"from Row 3 "
                f"({len(df)} data row(s))."
            )

    # ========================================================
    # SAVE UPDATED TEMPLATE
    # ========================================================

    output = BytesIO()

    try:

        target_wb.save(
            output
        )

    except Exception as e:

        raise ProcessingError(
            "Could not save the updated Template workbook: "
            f"{e}"
        )

    # ========================================================
    # RETURN
    # ========================================================

    output.seek(0)

    log_lines.append(
        "All sheets processed successfully."
    )

    return (
        output.read(),
        sheets,
        log_lines
    )